from django.contrib import messages
from django.shortcuts import redirect, render, get_object_or_404
from django.core.paginator import Paginator
from django.db.models import Q
from django.db.models import Prefetch
from decimal import Decimal
from django.http import JsonResponse
from django.utils import timezone
from django.http import HttpResponse
from operator import or_
from functools import reduce

import csv
from datetime import timedelta

from django.db.models import Sum, Count, DecimalField, Value
from django.utils.dateparse import parse_date
from django.db.models.functions import Coalesce, TruncMonth

from apps.accounts.models import UserAccount
from apps.dashboard.models import Party, PartyPaymentQR, PartyPurchase, PurchasePaymentInstallment, QuickLink, QuickLinkCategory, Notification, NotificationRecipient, IncomeSource, IncomePayment

from apps.dashboard.session import (
    get_logged_in_user,
    login_user,
    logout_user,
)

from apps.dashboard.decorators import (
    login_required,
    dashboard_access_required,
)

from django.shortcuts import render
from django.http import FileResponse

from apps.pdf_overlay.models import OverlayImage
from apps.pdf_overlay.services.pdf_processor import apply_overlay





# ==========================================================
# LOGIN
# ==========================================================
def login_view(request):

    # ------------------------------------------------------
    # STEP 1: CHECK ALREADY LOGGED-IN USER
    # ------------------------------------------------------

    current_user = get_logged_in_user(
        request
    )


    if current_user:

        # If already has dashboard access
        if current_user.has_dashboard_access:

            return redirect(
                "dashboard_view"
            )


    # ------------------------------------------------------
    # STEP 2: HANDLE POST
    # ------------------------------------------------------

    if request.method == "POST":

        # --------------------------------------------------
        # GET FORM DATA
        # --------------------------------------------------

        email = request.POST[
            "email"
        ].strip().lower()

        password = request.POST[
            "password"
        ]

        remember_me = request.POST.get(
            "remember_me"
        )


        # --------------------------------------------------
        # FIND USER
        # --------------------------------------------------

        try:

            user = (
                UserAccount.objects
                .select_related("role")
                .get(
                    email=email
                )
            )

        except UserAccount.DoesNotExist:

            return render(
                request,
                "dashboard/login.html",
                {
                    "error": (
                        "Invalid email or password."
                    ),
                    "email": email,
                }
            )


        # --------------------------------------------------
        # CHECK ACTIVE
        # --------------------------------------------------

        if not user.is_active:

            return render(
                request,
                "dashboard/login.html",
                {
                    "error": (
                        "Your account has been deactivated."
                    ),
                    "email": email,
                }
            )


        # --------------------------------------------------
        # CHECK PASSWORD
        # --------------------------------------------------

        if not user.check_password(
            password
        ):

            return render(
                request,
                "dashboard/login.html",
                {
                    "error": (
                        "Invalid email or password."
                    ),
                    "email": email,
                }
            )


        # --------------------------------------------------
        # CHECK DASHBOARD ACCESS
        # --------------------------------------------------

        if not user.has_dashboard_access:

            return render(
                request,
                "dashboard/login.html",
                {
                    "error": (
                        "You do not have permission "
                        "to access the dashboard."
                    ),
                    "email": email,
                }
            )


        # --------------------------------------------------
        # LOGIN
        # --------------------------------------------------

        login_user(
            request,
            user
        )


        # --------------------------------------------------
        # REMEMBER ME
        # --------------------------------------------------

        if remember_me:

            # 30 days
            request.session.set_expiry(
                60 * 60 * 24 * 30
            )

        else:

            # Browser session
            request.session.set_expiry(
                0
            )

        # --------------------------------------------------
        # REDIRECT DASHBOARD
        # --------------------------------------------------

        return redirect(
            "dashboard_view"
        )


    # ------------------------------------------------------
    # GET LOGIN PAGE
    # ------------------------------------------------------

    return render(
        request,
        "dashboard/login.html"
    )


# ==========================================================
# LOGOUT
# ==========================================================
@login_required
def logout_view(request):

    # ------------------------------------------------------
    # DESTROY CUSTOM SESSION
    # ------------------------------------------------------

    logout_user(
        request
    )


    # ------------------------------------------------------
    # SUCCESS MESSAGE
    # ------------------------------------------------------

    messages.success(
        request,
        "You have been logged out successfully."
    )


    # ------------------------------------------------------
    # REDIRECT LOGIN
    # ------------------------------------------------------

    return redirect(
        "login_view"
    )


# ==========================================================
# DASHBOARD
# ==========================================================
@dashboard_access_required
def dashboard_view(request):

    # =====================================================
    # CURRENT USER
    # =====================================================

    user = request.current_user

    # =====================================================
    # PARTY STATISTICS
    # =====================================================

    total_parties = Party.objects.count()

    active_parties = Party.objects.filter(
        is_active=True
    ).count()

    inactive_parties = Party.objects.filter(
        is_active=False
    ).count()

    gst_parties = (
        Party.objects
        .exclude(gst_number="")
        .exclude(gst_number__isnull=True)
        .count()
    )

    bank_parties = (
        Party.objects
        .exclude(bank_name="")
        .exclude(bank_name__isnull=True)
        .exclude(account_number="")
        .exclude(account_number__isnull=True)
        .count()
    )

    total_qrs = PartyPaymentQR.objects.count()

    active_qrs = PartyPaymentQR.objects.filter(
        is_active=True
    ).count()

    # =====================================================
    # PARTY TYPES
    # =====================================================

    party_type_data = []

    for value, label in Party.PartyType.choices:

        count = Party.objects.filter(
            party_type=value
        ).count()

        party_type_data.append(
            {
                "value": value,
                "label": label,
                "count": count,
            }
        )

    # =====================================================
    # RECENT PARTIES
    # =====================================================

    recent_parties = (
        Party.objects
        .order_by("-created_at")[:8]
    )

    # =====================================================
    # RECENTLY UPDATED PARTIES
    # =====================================================

    recently_updated_parties = (
        Party.objects
        .order_by("-updated_at")[:5]
    )

    # =====================================================
    # PARTIES NEEDING ATTENTION
    # =====================================================

    parties_without_gst = (
        Party.objects
        .filter(is_active=True)
        .filter(
            gst_number=""
        )
        .order_by("name")[:5]
    )

    parties_without_bank = (
        Party.objects
        .filter(is_active=True)
        .filter(
            bank_name=""
        )
        .order_by("name")[:5]
    )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {
        "user": user,

        # Statistics
        "total_parties": total_parties,
        "active_parties": active_parties,
        "inactive_parties": inactive_parties,
        "gst_parties": gst_parties,
        "bank_parties": bank_parties,
        "total_qrs": total_qrs,
        "active_qrs": active_qrs,

        # Party types
        "party_type_data": party_type_data,

        # Lists
        "recent_parties": recent_parties,
        "recently_updated_parties": recently_updated_parties,

        # Attention
        "parties_without_gst": parties_without_gst,
        "parties_without_bank": parties_without_bank,
    }

    return render(
        request,
        "dashboard/dashboard.html",
        context,
    )




# =========================================================
# SHARED PARTY QUERYSET
# =========================================================

def _get_filtered_party_queryset(request):
    """Build the same filtered queryset for listing and export."""

    search = request.GET.get("search", "").strip()[:150]
    party_type = request.GET.get("party_type", "").strip()
    status = request.GET.get("status", "").strip().lower()
    gst = request.GET.get("gst", "").strip().lower()
    bank = request.GET.get("bank", "").strip().lower()
    sort = request.GET.get("sort", "-pk").strip()

    sort_options = {
        "name": "Name: A–Z",
        "-name": "Name: Z–A",
        "party_type": "Party type: A–Z",
        "-party_type": "Party type: Z–A",
        "city": "City: A–Z",
        "-city": "City: Z–A",
        "state": "State: A–Z",
        "-state": "State: Z–A",
        "is_active": "Inactive first",
        "-is_active": "Active first",
        "pk": "Oldest added",
        "-pk": "Newest added",
    }

    if sort not in sort_options:
        sort = "-pk"

    valid_party_types = {
        str(value)
        for value, _label in Party.PartyType.choices
    }

    if party_type not in valid_party_types:
        party_type = ""

    if status not in {"", "active", "inactive"}:
        status = ""

    if gst not in {"", "available", "missing"}:
        gst = ""

    if bank not in {"", "available", "missing"}:
        bank = ""

    # Base queryset
    parties = Party.objects.all().prefetch_related("payment_qrs")

    # Search across the existing directory fields.
    if search:
        search_fields = (
            "name",
            "contact_person",
            "phone",
            "alternate_phone",
            "email",
            "gst_number",
            "city",
            "state",
            "pincode",
            "bank_name",
            "account_holder_name",
            "account_number",
            "ifsc_code",
            "branch_name",
            "upi_id",
            "gpay_number",
        )

        search_query = reduce(
            or_,
            [
                Q(**{f"{field}__icontains": search})
                for field in search_fields
            ],
        )

        parties = parties.filter(search_query)

    # Party type
    if party_type:
        parties = parties.filter(party_type=party_type)

    # Active / inactive
    if status == "active":
        parties = parties.filter(is_active=True)
    elif status == "inactive":
        parties = parties.filter(is_active=False)

    # GST availability
    if gst == "available":
        parties = (
            parties
            .exclude(gst_number__isnull=True)
            .exclude(gst_number__exact="")
        )

    elif gst == "missing":
        parties = parties.filter(
            Q(gst_number__isnull=True)
            | Q(gst_number__exact="")
        )

    # Bank / payment profile
    bank_fields = (
        "bank_name",
        "account_holder_name",
        "account_number",
        "ifsc_code",
        "branch_name",
        "upi_id",
        "gpay_number",
    )

    if bank in {"available", "missing"}:
        bank_query = reduce(
            or_,
            [
                Q(**{f"{field}__isnull": False})
                & ~Q(**{field: ""})
                for field in bank_fields
            ],
        )

        if bank == "available":
            parties = parties.filter(bank_query)
        else:
            parties = parties.exclude(bank_query)

    # Stable ordering
    parties = parties.order_by(sort, "pk")

    filters = {
        "search": search,
        "party_type": party_type,
        "status": status,
        "gst": gst,
        "bank": bank,
        "sort": sort,
    }

    return parties, filters, sort_options

# =========================================================
# PARTY DIRECTORY VIEW
# =========================================================

@dashboard_access_required
def party_list(request):
    parties, filters, sort_options = _get_filtered_party_queryset(request)

    # Page size
    page_size_options = (15, 25, 50, 100)

    try:
        per_page = int(request.GET.get("per_page", 15))
    except (TypeError, ValueError):
        per_page = 15

    if per_page not in page_size_options:
        per_page = 15

    # Global statistics across ALL parties
    global_stats = Party.objects.aggregate(
        total=Count("pk"),
        active=Count("pk", filter=Q(is_active=True)),
        inactive=Count("pk", filter=Q(is_active=False)),
        with_gst=Count(
            "pk",
            filter=(
                ~Q(gst_number__isnull=True)
                & ~Q(gst_number__exact="")
            ),
        ),
    )

    total_parties = global_stats["total"] or 0
    active_parties = global_stats["active"] or 0
    inactive_parties = global_stats["inactive"] or 0
    gst_parties = global_stats["with_gst"] or 0

    # Pagination applies ONLY to the directory display.
    paginator = Paginator(parties, per_page)

    page_obj = paginator.get_page(request.GET.get("page", 1))

    page_range = paginator.get_elided_page_range(
        number=page_obj.number,
        on_each_side=2,
        on_ends=1,
    )

    filtered_count = paginator.count

    # Party-type chart
    type_totals = dict(
        Party.objects
        .values("party_type")
        .annotate(total=Count("pk"))
        .values_list("party_type", "total")
    )

    party_type_labels = dict(Party.PartyType.choices)

    party_type_chart = {
        "labels": [
            party_type_labels.get(value, str(value or "Unspecified"))
            for value in party_type_labels
            if type_totals.get(value, 0)
        ],
        "values": [
            type_totals.get(value, 0)
            for value in party_type_labels
            if type_totals.get(value, 0)
        ],
    }

    # State chart
    state_rows = list(
        Party.objects
        .exclude(state__isnull=True)
        .exclude(state__exact="")
        .values("state")
        .annotate(total=Count("pk"))
        .order_by("-total", "state")[:8]
    )

    state_chart = {
        "labels": [row["state"] for row in state_rows],
        "values": [row["total"] for row in state_rows],
    }

    status_chart = {
        "labels": ["Active", "Inactive"],
        "values": [active_parties, inactive_parties],
    }

    # Preserve query parameters for pagination.
    query_params = request.GET.copy()
    query_params.pop("page", None)
    pagination_query = query_params.urlencode()

    context = {
        "page_obj": page_obj,
        "parties": page_obj.object_list,
        "page_range": page_range,

        # Filters
        **filters,

        # Options
        "party_types": Party.PartyType.choices,
        "page_size_options": page_size_options,
        "sort_options": sort_options.items(),
        "per_page": per_page,

        # Global statistics
        "total_parties": total_parties,
        "active_parties": active_parties,
        "inactive_parties": inactive_parties,
        "gst_parties": gst_parties,

        # Filtered statistics
        "filtered_parties_count": filtered_count,
        "has_active_filters": bool(
            filters["search"]
            or filters["party_type"]
            or filters["status"]
            or filters["gst"]
            or filters["bank"]
        ),
        "pagination_query": pagination_query,

        # Charts
        "status_chart": status_chart,
        "party_type_chart": party_type_chart,
        "state_chart": state_chart,

        # Percentages
        "active_percentage": round(
            active_parties / total_parties * 100, 1
        ) if total_parties else 0,

        "gst_percentage": round(
            gst_parties / total_parties * 100, 1
        ) if total_parties else 0,
    }

    is_ajax = (
        request.headers.get("X-Requested-With")
        == "XMLHttpRequest"
    )

    is_htmx = (
        request.headers.get("HX-Request", "").lower()
        == "true"
    )

    if is_ajax or is_htmx:
        return render(
            request,
            "dashboard/party_table.html",
            context,
        )

    return render(
        request,
        "dashboard/party_list.html",
        context,
    )

# =========================================================
# EXCEL EXPORT HELPERS
# =========================================================

def _get_export_columns(model):
    columns = []

    for field in model._meta.concrete_fields:
        label = str(field.verbose_name).replace("_", " ").title()

        if field.is_relation and (
            field.many_to_one or field.one_to_one
        ):
            columns.append((field, "id", f"{label} ID"))
            columns.append((field, "related", f"{label} Display"))

        elif field.choices:
            columns.append((field, "value", f"{label} Value"))
            columns.append((field, "display", f"{label} Display"))

        else:
            columns.append((field, "value", label))

    return columns


def _get_export_value(record, field, value_type):
    if value_type == "id":
        return getattr(record, field.attname, None)

    if value_type == "related":
        related_object = getattr(record, field.name, None)
        return str(related_object) if related_object is not None else ""

    if value_type == "display":
        display_method = getattr(
            record,
            f"get_{field.name}_display",
            None,
        )

        if callable(display_method):
            return display_method()

    return getattr(record, field.attname, None)


def _excel_safe_value(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.make_naive(value)
        return value

    if isinstance(value, date):
        return value

    if isinstance(value, bool):
        return value

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, (int, float)):
        return value

    if isinstance(value, bytes):
        value = value.decode("utf-8", errors="replace")

    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)

    text_value = str(value)

    # Protect text fields from being interpreted as Excel formulas.
    if text_value.startswith(("=", "+", "-", "@")):
        text_value = "'" + text_value

    return text_value


def _style_excel_sheet(worksheet):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="153030",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 32
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    if worksheet.max_row and worksheet.max_column:
        worksheet.auto_filter.ref = worksheet.dimensions

    sample_end = min(worksheet.max_row, 250)

    for column_index in range(1, worksheet.max_column + 1):
        max_length = 12

        for row_index in range(1, sample_end + 1):
            cell_value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            max_length = max(
                max_length,
                len(str(cell_value or "")),
            )

        letter = get_column_letter(column_index)

        worksheet.column_dimensions[letter].width = min(
            max_length + 2,
            40,
        )

# =========================================================
# EXPORT ALL PARTIES TO EXCEL
# =========================================================

@dashboard_access_required
def export_parties_excel(request):
    # Use the same filters as the directory.
    parties, filters, _sort_options = _get_filtered_party_queryset(request)

    # Optimise foreign-key display values.
    related_fields = [
        field.name
        for field in Party._meta.concrete_fields
        if field.is_relation
        and (field.many_to_one or field.one_to_one)
    ]

    if related_fields:
        parties = parties.select_related(*related_fields)

    # Detect the related QR model if the relationship exists.
    qr_model = None

    try:
        qr_model = Party._meta.get_field("payment_qrs").related_model
    except (FieldDoesNotExist, AttributeError):
        pass

    # Create workbook and primary Party worksheet.
    workbook = Workbook()
    party_sheet = workbook.active
    party_sheet.title = "Parties"

    party_columns = _get_export_columns(Party)

    party_sheet.append([
        label
        for _field, _value_type, label in party_columns
    ])

    # Optional related QR worksheet.
    qr_sheet = None
    qr_columns = []

    if qr_model is not None:
        qr_sheet = workbook.create_sheet("Payment QRs")
        qr_columns = _get_export_columns(qr_model)

        qr_sheet.append(
            ["Party ID", "Party Name"]
            + [
                label
                for _field, _value_type, label in qr_columns
            ]
        )

    # IMPORTANT:
    # This is the complete filtered queryset.
    # No Paginator, page number, or page-size restriction.
    for party in parties.iterator(chunk_size=500):
        party_sheet.append([
            _excel_safe_value(
                _get_export_value(party, field, value_type)
            )
            for field, value_type, _label in party_columns
        ])

        # Export related QR data into its own worksheet.
        if qr_sheet is not None:
            qr_manager = getattr(party, "payment_qrs", None)

            if qr_manager is not None:
                for qr_record in qr_manager.all():
                    qr_sheet.append(
                        [
                            _excel_safe_value(party.pk),
                            _excel_safe_value(party.name),
                        ]
                        + [
                            _excel_safe_value(
                                _get_export_value(
                                    qr_record,
                                    field,
                                    value_type,
                                )
                            )
                            for field, value_type, _label in qr_columns
                        ]
                    )

    _style_excel_sheet(party_sheet)

    if qr_sheet is not None:
        _style_excel_sheet(qr_sheet)

    # Download as a real .xlsx Excel file.
    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    filename = f"parties_export_{timestamp}.xlsx"

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    workbook.save(response)

    return response

@dashboard_access_required
def party_detail(request, pk):
    party = get_object_or_404(
        Party.objects.prefetch_related("payment_qrs"),
        pk=pk
    )

    return render(
        request,
        "dashboard/party_detail.html",
        {
            "party": party,
        }
    )


def party_orders_view(request, party_id):
    """
    Display all purchase orders for a particular party.
    """

    party = get_object_or_404(
        Party,
        pk=party_id,
    )


    orders = (
        PartyPurchase.objects
        .filter(party=party)
        .select_related("party")
        .order_by(
            "-purchase_date",
            "-created_at",
        )
    )

    summary = orders.aggregate(
        total_purchase=Sum("grand_total"),
        total_paid=Sum("paid_amount"),
        total_due=Sum("due_amount"),
    )

    total_purchase = (
        summary["total_purchase"]
        or Decimal("0.00")
    )

    total_paid = (
        summary["total_paid"]
        or Decimal("0.00")
    )

    total_due = (
        summary["total_due"]
        or Decimal("0.00")
    )

    context = {
        "party": party,
        "orders": orders,

        "total_orders": orders.count(),

        "total_purchase": total_purchase,

        "total_paid": total_paid,

        "total_due": total_due,
    }

    return render(
        request,
        "dashboard/party_orders.html",
        context,
    )

def party_purchase_detail_view(request, purchase_id):

    purchase = get_object_or_404(
        PartyPurchase.objects.select_related("party"),
        pk=purchase_id,
    )

    installments = (
        PurchasePaymentInstallment.objects
        .filter(purchase=purchase)
        .select_related("payment_method")
        .order_by("installment_number")
    )

    installment_summary = installments.aggregate(
        total_installment=Sum("installment_amount"),
        total_paid=Sum("paid_amount"),
        total_remaining=Sum("remaining_amount"),
    )

    context = {
        "purchase": purchase,
        "party": purchase.party,
        "installments": installments,

        "total_installments": installments.count(),

        "total_installment": (
            installment_summary["total_installment"]
            or Decimal("0.00")
        ),

        "installment_paid": (
            installment_summary["total_paid"]
            or Decimal("0.00")
        ),

        "installment_remaining": (
            installment_summary["total_remaining"]
            or Decimal("0.00")
        ),
    }

    return render(
        request,
        "dashboard/party_purchase_detail.html",
        context,
    )
def purchase_installments_view(request, purchase_id):

    purchase = get_object_or_404(
        PartyPurchase.objects.select_related("party"),
        pk=purchase_id,
    )

    installments = (
        PurchasePaymentInstallment.objects
        .filter(purchase=purchase)
        .select_related("payment_method")
        .order_by("installment_number")
    )

    summary = installments.aggregate(
        total_installment=Sum("installment_amount"),
        total_paid=Sum("paid_amount"),
        total_remaining=Sum("remaining_amount"),
    )

    context = {
        "purchase": purchase,
        "party": purchase.party,
        "installments": installments,

        "total_installments": installments.count(),

        "total_installment": (
            summary["total_installment"]
            or Decimal("0.00")
        ),

        "total_paid": (
            summary["total_paid"]
            or Decimal("0.00")
        ),

        "total_remaining": (
            summary["total_remaining"]
            or Decimal("0.00")
        ),
    }

    return render(
        request,
        "dashboard/purchase_installments.html",
        context,
    )



import calendar
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from django.core.paginator import Paginator
from django.core.exceptions import ValidationError
from django.db.models import Avg, Count, DecimalField, Q, Sum, Value
from django.db.models.functions import Coalesce, TruncDay, TruncMonth
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.db import models
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.dashboard.decorators import dashboard_access_required


MONEY_FIELD = DecimalField(max_digits=16, decimal_places=2)
ZERO = Value(Decimal("0.00"), output_field=MONEY_FIELD)

PERIOD_OPTIONS = [
    ("", "All time"),
    ("today", "Today"),
    ("week", "Last 7 days"),
    ("month", "This month"),
    ("quarter", "This quarter"),
    ("year", "This year"),
    ("last_30", "Last 30 days"),
    ("last_90", "Last 90 days"),
    ("previous_month", "Previous month"),
    ("previous_quarter", "Previous quarter"),
    ("custom", "Custom dates"),
]


def _amount_total(queryset, field_name):
    return queryset.aggregate(
        total=Coalesce(Sum(field_name), ZERO, output_field=MONEY_FIELD)
    )["total"] or Decimal("0.00")


def _percentage(numerator, denominator):
    numerator = Decimal(str(numerator or 0))
    denominator = Decimal(str(denominator or 0))
    if denominator <= 0:
        return 0.0
    return round(float(numerator / denominator * Decimal("100")), 1)


def _first_date_field(model, candidates):
    """Find the first real DateField/DateTimeField from an ordered list."""
    for name in candidates:
        try:
            field = model._meta.get_field(name)
        except Exception:
            continue
        if isinstance(field, (models.DateField, models.DateTimeField)):
            return name
    return None


def _all_field_names(model):
    names = set()
    for field in model._meta.get_fields():
        if getattr(field, "concrete", False):
            names.add(field.name)
    return names


def _month_start(value):
    return value.replace(day=1)


def _previous_month_range(today):
    first_this_month = today.replace(day=1)
    last_previous_month = first_this_month - timedelta(days=1)
    return last_previous_month.replace(day=1), last_previous_month


def _quarter_start(value):
    month = ((value.month - 1) // 3) * 3 + 1
    return date(value.year, month, 1)


def _previous_quarter_range(today):
    start_this_quarter = _quarter_start(today)
    end_previous_quarter = start_this_quarter - timedelta(days=1)
    return _quarter_start(end_previous_quarter), end_previous_quarter


def _date_range_from_request(request, today=None):
    """Return (period, start_date, end_date, label, validation_message)."""
    today = today or timezone.localdate()
    raw_period = (request.GET.get("period") or "").strip()
    valid_periods = {item[0] for item in PERIOD_OPTIONS}
    period = raw_period if raw_period in valid_periods else ""
    start_date = parse_date((request.GET.get("start_date") or "").strip())
    end_date = parse_date((request.GET.get("end_date") or "").strip())
    message = ""

    if period == "today":
        start_date = end_date = today
    elif period == "week":
        start_date, end_date = today - timedelta(days=6), today
    elif period == "month":
        start_date, end_date = _month_start(today), today
    elif period == "quarter":
        start_date, end_date = _quarter_start(today), today
    elif period == "year":
        start_date, end_date = date(today.year, 1, 1), today
    elif period == "last_30":
        start_date, end_date = today - timedelta(days=29), today
    elif period == "last_90":
        start_date, end_date = today - timedelta(days=89), today
    elif period == "previous_month":
        start_date, end_date = _previous_month_range(today)
    elif period == "previous_quarter":
        start_date, end_date = _previous_quarter_range(today)
    elif period == "custom":
        pass
    elif start_date or end_date:
        period = "custom"
    else:
        start_date = end_date = None

    if start_date and end_date and start_date > end_date:
        start_date, end_date = end_date, start_date
        period = "custom"
        message = "Start and end dates were reversed, so the range has been corrected."

    if start_date and end_date:
        if start_date == end_date:
            label = start_date.strftime("%d %b %Y")
        else:
            label = f"{start_date:%d %b %Y} – {end_date:%d %b %Y}"
    elif start_date:
        label = f"From {start_date:%d %b %Y}"
    elif end_date:
        label = f"Through {end_date:%d %b %Y}"
    else:
        label = "All time"

    return period, start_date, end_date, label, message


def _filter_dates(queryset, model, field_name, start_date, end_date):
    if not field_name or not (start_date or end_date):
        return queryset
    try:
        field = model._meta.get_field(field_name)
    except Exception:
        return queryset

    lookup_base = f"{field_name}__date" if isinstance(field, models.DateTimeField) else field_name
    filters = {}
    if start_date:
        filters[f"{lookup_base}__gte"] = start_date
    if end_date:
        filters[f"{lookup_base}__lte"] = end_date
    return queryset.filter(**filters)


def _querystring_without(request, *keys):
    params = request.GET.copy()
    for key in keys:
        params.pop(key, None)
    return params.urlencode()


def _add_filter_context(context, request, period, start_date, end_date, label, message=""):
    context.update({
        "period_options": PERIOD_OPTIONS,
        "selected_period": period,
        "start_date": start_date.isoformat() if start_date else "",
        "end_date": end_date.isoformat() if end_date else "",
        "period_label": label,
        "filter_message": message,
        "pagination_query": _querystring_without(request, "page", "export"),
    })
    return context


def _xlsx_response(filename, headers, rows):
    """Create a formatted .xlsx file that opens in Microsoft Excel/LibreOffice."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(list(headers))

    header_fill = PatternFill(fill_type="solid", fgColor="153030")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 26

    for row in rows:
        cleaned = []
        for value in row:
            if isinstance(value, datetime) and timezone.is_aware(value):
                value = timezone.make_naive(value)
            if isinstance(value, UUID):
                value = str(value)
            cleaned.append(value)
        sheet.append(cleaned)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    currency_columns = []
    for idx, header in enumerate(headers, start=1):
        normalized = str(header).lower()
        if any(word in normalized for word in ("amount", "grand total", "purchase value", "paid", "due")):
            currency_columns.append(idx)
    for column_index in currency_columns:
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row=row_index, column=column_index).number_format = '"₹"#,##0.00;[Red]-"₹"#,##0.00'
    for column_cells in sheet.columns:
        values = [str(cell.value or "") for cell in column_cells[:250]]
        width = min(max(max((len(value) for value in values), default=10) + 2, 12), 42)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _trend_data(queryset, model, date_field_name, amount_field, start_date, end_date):
    if not date_field_name:
        return {}, "month"
    span = (end_date - start_date).days if start_date and end_date else 3650
    grouping = "day" if span <= 62 else "month"
    trunc = TruncDay if grouping == "day" else TruncMonth
    rows = (
        _filter_dates(queryset, model, date_field_name, start_date, end_date)
        .order_by()
        .annotate(bucket=trunc(date_field_name))
        .values("bucket")
        .annotate(total=Coalesce(Sum(amount_field), ZERO, output_field=MONEY_FIELD))
        .order_by("bucket")
    )
    result = {}
    for row in rows:
        bucket = row["bucket"]
        if bucket is None:
            continue
        if isinstance(bucket, datetime):
            bucket = timezone.localtime(bucket).date() if timezone.is_aware(bucket) else bucket.date()
        result[bucket] = float(row["total"] or 0)
    return result, grouping


def _trend_label(bucket, grouping):
    if grouping == "day":
        return bucket.strftime("%d %b")
    return bucket.strftime("%b %Y")


def _searchable_purchase_q(model, search):
    """Build safe search Q objects only for actual text fields in this model."""
    query = Q()
    found = False
    for field in model._meta.fields:
        if isinstance(field, (models.CharField, models.TextField)):
            query |= Q(**{f"{field.name}__icontains": search})
            found = True
        elif field.is_relation and field.related_model:
            related = field.related_model
            for related_field in related._meta.fields:
                if isinstance(related_field, (models.CharField, models.TextField)):
                    query |= Q(**{f"{field.name}__{related_field.name}__icontains": search})
                    found = True
    return query if found else None


def _party_field_info(model):
    try:
        field = model._meta.get_field("party")
        if field.is_relation and field.related_model:
            return field, field.related_model
    except Exception:
        pass
    return None, None


def _purchase_field_values(purchase, date_field):
    purchase_date = getattr(purchase, date_field, None) if date_field else None
    number = getattr(purchase, "purchase_number", None) or str(purchase.pk)
    party = getattr(purchase, "party", None)
    grand_total = getattr(purchase, "grand_total", Decimal("0.00")) or Decimal("0.00")
    paid_amount = getattr(purchase, "paid_amount", Decimal("0.00")) or Decimal("0.00")
    due_amount = getattr(purchase, "due_amount", Decimal("0.00")) or Decimal("0.00")
    status_field = "status" if "status" in _all_field_names(purchase.__class__) else None
    status_value = getattr(purchase, status_field, "") if status_field else ""
    status_label = purchase.get_status_display() if status_field and hasattr(purchase, "get_status_display") else status_value
    return purchase_date, number, party, grand_total, paid_amount, due_amount, status_value, status_label


# ---------------------------------------------------------------------------
# 1. FINANCE DASHBOARD — numbers and charts only; no transaction records.
# ---------------------------------------------------------------------------
@dashboard_access_required
def finance_dashboard(request):
    today = timezone.localdate()
    period, start_date, end_date, period_label, filter_message = _date_range_from_request(request, today)

    income_date_field = _first_date_field(IncomePayment, ["payment_date", "created_at", "created_on"])
    purchase_date_field = _first_date_field(PartyPurchase, ["purchase_date", "order_date", "date", "created_at", "created_on", "created"])
    installment_date_field = _first_date_field(PurchasePaymentInstallment, ["payment_date", "due_date", "created_at", "created_on"])

    income_all = IncomePayment.objects.select_related("income_source")
    income_period = _filter_dates(income_all, IncomePayment, income_date_field, start_date, end_date)
    completed = income_period.filter(payment_status=IncomePayment.PaymentStatus.COMPLETED)
    pending = income_period.filter(payment_status=IncomePayment.PaymentStatus.PENDING)
    failed = income_period.filter(payment_status=IncomePayment.PaymentStatus.FAILED)
    cancelled = income_period.filter(payment_status=IncomePayment.PaymentStatus.CANCELLED)

    purchases_all = PartyPurchase.objects.all()
    purchases_period = _filter_dates(purchases_all, PartyPurchase, purchase_date_field, start_date, end_date)
    installments_all = PurchasePaymentInstallment.objects.select_related("purchase", "purchase__party", "payment_method")
    installments_period = _filter_dates(installments_all, PurchasePaymentInstallment, installment_date_field, start_date, end_date)

    income_stats = {
        "total_count": income_period.count(),
        "total_amount": _amount_total(income_period, "amount"),
        "completed_count": completed.count(),
        "completed_amount": _amount_total(completed, "amount"),
        "pending_count": pending.count(),
        "pending_amount": _amount_total(pending, "amount"),
        "failed_count": failed.count(),
        "failed_amount": _amount_total(failed, "amount"),
        "cancelled_count": cancelled.count(),
        "cancelled_amount": _amount_total(cancelled, "amount"),
    }
    income_stats["completion_percentage"] = _percentage(income_stats["completed_count"], income_stats["total_count"])

    purchase_stats = {
        "purchase_count": purchases_period.count(),
        "total_purchase": _amount_total(purchases_period, "grand_total"),
        "purchase_paid": _amount_total(purchases_period, "paid_amount"),
        "purchase_due": _amount_total(purchases_period, "due_amount"),
    }
    purchase_stats["collection_percentage"] = _percentage(purchase_stats["purchase_paid"], purchase_stats["total_purchase"])

    installment_stats = {
        "installment_count": installments_period.count(),
        "scheduled_amount": _amount_total(installments_period, "installment_amount"),
        "paid_amount": _amount_total(installments_period, "paid_amount"),
        "remaining_amount": _amount_total(installments_period, "remaining_amount"),
    }
    installment_stats["collection_percentage"] = _percentage(installment_stats["paid_amount"], installment_stats["scheduled_amount"])

    # Overdue is an all-open-items risk indicator, independent of the selected period.
    installment_fields = _all_field_names(PurchasePaymentInstallment)
    if "due_date" in installment_fields:
        overdue = installments_all.filter(due_date__lt=today, remaining_amount__gt=0)
        installment_stats["overdue_count"] = overdue.count()
        installment_stats["overdue_amount"] = _amount_total(overdue, "remaining_amount")
    else:
        installment_stats["overdue_count"] = 0
        installment_stats["overdue_amount"] = Decimal("0.00")

    income_trend, grouping = _trend_data(completed, IncomePayment, income_date_field, "amount", start_date, end_date)
    purchase_trend, purchase_grouping = _trend_data(purchases_period, PartyPurchase, purchase_date_field, "grand_total", start_date, end_date)
    # Both series use the same grouping rule so the combined chart aligns.
    if grouping != purchase_grouping:
        grouping = "month" if grouping == "month" or purchase_grouping == "month" else "day"
        income_trend, _ = _trend_data(completed, IncomePayment, income_date_field, "amount", start_date, end_date)
        purchase_trend, _ = _trend_data(purchases_period, PartyPurchase, purchase_date_field, "grand_total", start_date, end_date)
    buckets = sorted(set(income_trend) | set(purchase_trend))
    trend_chart = {
        "labels": [_trend_label(bucket, grouping) for bucket in buckets],
        "income": [income_trend.get(bucket, 0) for bucket in buckets],
        "purchases": [purchase_trend.get(bucket, 0) for bucket in buckets],
    }

    source_rows = list(
        completed.values("income_source__name")
        .annotate(total=Coalesce(Sum("amount"), ZERO, output_field=MONEY_FIELD), count=Count("pk"))
        .order_by("-total")
    )
    source_chart = {
        "labels": [row["income_source__name"] or "Unassigned" for row in source_rows],
        "values": [float(row["total"] or 0) for row in source_rows],
    }

    status_choices = list(IncomePayment.PaymentStatus.choices)
    status_chart = {
        "labels": [label for value, label in status_choices],
        "values": [income_period.filter(payment_status=value).count() for value, label in status_choices],
    }

    method_rows = list(
        completed.values("payment_method")
        .annotate(total=Coalesce(Sum("amount"), ZERO, output_field=MONEY_FIELD), count=Count("pk"))
        .order_by("-total")
    )
    method_labels = dict(IncomePayment.PaymentMethod.choices)
    method_chart = {
        "labels": [method_labels.get(row["payment_method"], row["payment_method"] or "Unspecified") for row in method_rows],
        "values": [float(row["total"] or 0) for row in method_rows],
    }

    installment_methods = list(
        installments_period.filter(paid_amount__gt=0)
        .values("payment_method__name")
        .annotate(total=Coalesce(Sum("paid_amount"), ZERO, output_field=MONEY_FIELD))
        .order_by("-total")
    )
    installment_method_chart = {
        "labels": [row["payment_method__name"] or "Unspecified" for row in installment_methods],
        "values": [float(row["total"] or 0) for row in installment_methods],
    }

    income_received = income_stats["completed_amount"]
    estimated_cash_movement = income_received - purchase_stats["purchase_paid"]
    leading_source = source_rows[0] if source_rows else None
    smart_insights = []
    if installment_stats["overdue_count"]:
        smart_insights.append({"type": "warning", "title": "Overdue installments", "message": f'{installment_stats["overdue_count"]} open installment(s) have ₹{installment_stats["overdue_amount"]:,.2f} outstanding.'})
    if income_stats["pending_count"]:
        smart_insights.append({"type": "info", "title": "Pending income", "message": f'{income_stats["pending_count"]} transaction(s), totalling ₹{income_stats["pending_amount"]:,.2f}, are pending in this period.'})
    if estimated_cash_movement < 0:
        smart_insights.append({"type": "warning", "title": "Cash movement needs attention", "message": "Recorded purchase payments exceed completed income for the selected period. This is an operational indicator, not net profit."})
    if leading_source:
        smart_insights.append({"type": "success", "title": "Leading income source", "message": f'{leading_source["income_source__name"] or "Unassigned"} is the leading source by completed income in this period.'})
    if not smart_insights:
        smart_insights.append({"type": "info", "title": "No immediate alerts", "message": "No pending income or overdue installment alerts were found in the available data."})

    context = {
        "today": today,
        "financial_stats": {
            "total_income_received": income_received,
            "total_income_recorded": income_stats["total_amount"],
            "total_purchase_value": purchase_stats["total_purchase"],
            "total_purchase_paid": purchase_stats["purchase_paid"],
            "total_purchase_due": purchase_stats["purchase_due"],
            "estimated_cash_surplus": estimated_cash_movement,
            "estimated_margin_percentage": _percentage(income_received - purchase_stats["total_purchase"], income_received),
        },
        "income_stats": income_stats,
        "purchase_stats": purchase_stats,
        "installment_stats": installment_stats,
        "smart_insights": smart_insights,
        "trend_chart": trend_chart,
        "source_chart": source_chart,
        "status_chart": status_chart,
        "method_chart": method_chart,
        "installment_method_chart": installment_method_chart,
        "purchase_date_filter_supported": bool(purchase_date_field),
        "purchase_date_field": purchase_date_field or "",
        "income_date_field": income_date_field or "",
    }
    _add_filter_context(context, request, period, start_date, end_date, period_label, filter_message)
    return render(request, "dashboard/finance_dashboard.html", context)


# ---------------------------------------------------------------------------
# 2. INCOME TRANSACTIONS — filtered/paginated records plus Excel export.
# ---------------------------------------------------------------------------
@dashboard_access_required
def income_transactions(request):
    today = timezone.localdate()
    period, start_date, end_date, period_label, filter_message = _date_range_from_request(request, today)
    search = (request.GET.get("q") or "").strip()
    status = (request.GET.get("status") or "").strip()
    method = (request.GET.get("method") or "").strip()
    source = (request.GET.get("source") or "").strip()
    date_field = _first_date_field(IncomePayment, ["payment_date", "created_at", "created_on"])

    payments = IncomePayment.objects.select_related("income_source")
    if search:
        payments = payments.filter(
            Q(income_source__name__icontains=search)
            | Q(transaction_id__icontains=search)
            | Q(description__icontains=search)
        )
    valid_statuses = dict(IncomePayment.PaymentStatus.choices)
    if status in valid_statuses:
        payments = payments.filter(payment_status=status)
    else:
        status = ""
    valid_methods = dict(IncomePayment.PaymentMethod.choices)
    if method in valid_methods:
        payments = payments.filter(payment_method=method)
    else:
        method = ""
    if source:
        try:
            source_exists = IncomeSource.objects.filter(pk=source).exists()
        except (ValueError, TypeError, ValidationError):
            source_exists = False
        if source_exists:
            payments = payments.filter(income_source_id=source)
        else:
            source = ""
    payments = _filter_dates(payments, IncomePayment, date_field, start_date, end_date).order_by("-" + date_field if date_field else "-pk", "-pk")

    stats = {
        "count": payments.count(),
        "amount": _amount_total(payments, "amount"),
        "completed_count": payments.filter(payment_status=IncomePayment.PaymentStatus.COMPLETED).count(),
        "completed_amount": _amount_total(payments.filter(payment_status=IncomePayment.PaymentStatus.COMPLETED), "amount"),
        "pending_count": payments.filter(payment_status=IncomePayment.PaymentStatus.PENDING).count(),
        "pending_amount": _amount_total(payments.filter(payment_status=IncomePayment.PaymentStatus.PENDING), "amount"),
    }

    if request.GET.get("export") == "xlsx":
        def rows():
            for payment in payments.iterator(chunk_size=1000):
                yield [
                    payment.payment_date,
                    payment.income_source.name if payment.income_source else "",
                    float(payment.amount or 0),
                    payment.get_payment_method_display() if hasattr(payment, "get_payment_method_display") else getattr(payment, "payment_method", ""),
                    getattr(payment, "transaction_id", ""),
                    payment.get_payment_status_display() if hasattr(payment, "get_payment_status_display") else getattr(payment, "payment_status", ""),
                    getattr(payment, "description", ""),
                    payment.pk,
                ]
        return _xlsx_response("yaarlynx_income_transactions.xlsx", ["Payment Date", "Income Source", "Amount", "Payment Method", "Transaction ID", "Status", "Description", "Record ID"], rows())

    page_obj = Paginator(payments, 25).get_page(request.GET.get("page"))
    context = {
        "payments": page_obj.object_list,
        "page_obj": page_obj,
        "stats": stats,
        "search_query": search,
        "selected_status": status,
        "selected_method": method,
        "selected_source": source,
        "payment_statuses": IncomePayment.PaymentStatus.choices,
        "payment_methods": IncomePayment.PaymentMethod.choices,
        "income_sources": IncomeSource.objects.filter(is_active=True).order_by("name") if "is_active" in _all_field_names(IncomeSource) else IncomeSource.objects.all().order_by("name"),
        "date_field": date_field or "",
    }
    _add_filter_context(context, request, period, start_date, end_date, period_label, filter_message)
    return render(request, "dashboard/income_transactions.html", context)


# ---------------------------------------------------------------------------
# 3. PURCHASE TRANSACTIONS — filtered/paginated records plus Excel export.
# ---------------------------------------------------------------------------
@dashboard_access_required
def purchase_transactions(request):
    today = timezone.localdate()
    period, start_date, end_date, period_label, filter_message = _date_range_from_request(request, today)
    search = (request.GET.get("q") or "").strip()
    selected_party = (request.GET.get("party") or "").strip()
    selected_status = (request.GET.get("status") or "").strip()
    date_field = _first_date_field(PartyPurchase, ["purchase_date", "order_date", "date", "created_at", "created_on", "created"])
    purchase_fields = _all_field_names(PartyPurchase)

    party_field, party_model = _party_field_info(PartyPurchase)
    purchases = PartyPurchase.objects.all()
    if party_field:
        purchases = purchases.select_related("party")
    if search:
        search_q = _searchable_purchase_q(PartyPurchase, search)
        if search_q is not None:
            purchases = purchases.filter(search_q)
    if selected_party and party_field:
        try:
            party_exists = party_model.objects.filter(pk=selected_party).exists()
        except (ValueError, TypeError, ValidationError):
            party_exists = False
        if party_exists:
            purchases = purchases.filter(party_id=selected_party)
        else:
            selected_party = ""

    status_field = None
    try:
        candidate = PartyPurchase._meta.get_field("status")
        if getattr(candidate, "choices", None):
            status_field = "status"
    except Exception:
        pass
    purchase_statuses = list(PartyPurchase._meta.get_field(status_field).choices) if status_field else []
    valid_purchase_statuses = dict(purchase_statuses)
    if status_field and selected_status in valid_purchase_statuses:
        purchases = purchases.filter(**{status_field: selected_status})
    else:
        selected_status = ""

    purchases = _filter_dates(purchases, PartyPurchase, date_field, start_date, end_date)
    if date_field:
        purchases = purchases.order_by("-" + date_field, "-pk")
    else:
        purchases = purchases.order_by("-pk")

    stats = {
        "count": purchases.count(),
        "value": _amount_total(purchases, "grand_total"),
        "paid": _amount_total(purchases, "paid_amount"),
        "due": _amount_total(purchases, "due_amount"),
    }
    stats["collection_percentage"] = _percentage(stats["paid"], stats["value"])

    # If export is selected, export the complete filtered queryset, not just the visible page.
    if request.GET.get("export") == "xlsx":
        def rows():
            for purchase in purchases.iterator(chunk_size=1000):
                values = _purchase_field_values(purchase, date_field)
                p_date, number, party, grand_total, paid_amount, due_amount, status_value, status_label = values
                yield [p_date, number, str(party) if party else "", float(grand_total), float(paid_amount), float(due_amount), status_label, purchase.pk]
        return _xlsx_response("yaarlynx_purchase_transactions.xlsx", ["Purchase Date", "Purchase Number", "Party / Supplier", "Grand Total", "Paid Amount", "Due Amount", "Status", "Record ID"], rows())

    page_obj = Paginator(purchases, 25).get_page(request.GET.get("page"))
    for purchase in page_obj.object_list:
        values = _purchase_field_values(purchase, date_field)
        purchase.finance_date = values[0]
        purchase.finance_number = values[1]
        purchase.finance_party = values[2]
        purchase.finance_grand_total = values[3]
        purchase.finance_paid_amount = values[4]
        purchase.finance_due_amount = values[5]
        purchase.finance_status_label = values[7]

    party_choices = []
    if party_model:
        party_choices = list(party_model.objects.all().order_by("pk"))

    context = {
        "purchases": page_obj.object_list,
        "page_obj": page_obj,
        "stats": stats,
        "search_query": search,
        "selected_party": selected_party,
        "selected_status": selected_status,
        "parties": party_choices,
        "purchase_statuses": purchase_statuses,
        "date_field": date_field or "",
        "party_filter_supported": bool(party_field),
    }
    _add_filter_context(context, request, period, start_date, end_date, period_label, filter_message)
    return render(request, "dashboard/purchase_transactions.html", context)


# ==========================================================
# USER PROFILE
# ==========================================================
@login_required
def profile_view(request):

    # ------------------------------------------------------
    # GET CURRENT LOGGED-IN USER
    # ------------------------------------------------------

    user = request.current_user


    # ------------------------------------------------------
    # RENDER PROFILE PAGE
    # ------------------------------------------------------

    return render(
        request,
        "dashboard/profile.html",
        {
            "user": user,
        }
    )

def quick_links(request):

    active_quick_links = (
        QuickLink.objects
        .filter(is_active=True)
        .order_by("display_order", "name")
    )

    categories = (
        QuickLinkCategory.objects
        .filter(is_active=True)
        .prefetch_related(
            Prefetch(
                "quick_links",
                queryset=active_quick_links,
                to_attr="active_quick_links"
            )
        )
        .order_by("display_order", "name")
    )

    # Remove categories that have no active quick links
    categories = [
        category
        for category in categories
        if category.active_quick_links
    ]

    return render(
        request,
        "dashboard/quick_links.html",
        {
            "categories": categories,
        }
    )

def notification_list(request):
    """
    Display notifications available for the logged-in user.
    """

    user_id = request.session.get("yaarlynx_user_id")

    if not user_id:
        return redirect("login")

    user = get_object_or_404(
        UserAccount,
        id=user_id
    )

    recipients = (
        NotificationRecipient.objects
        .filter(
            user=user,
            notification__is_active=True,
            notification__publish_at__lte=timezone.now(),
        )
        .select_related("notification")
        .order_by(
            "-notification__is_pinned",
            "-notification__publish_at",
            "-created_at",
        )
    )

    # Remove expired notifications
    recipients = [
        recipient
        for recipient in recipients
        if recipient.notification.is_published
    ]

    total_notifications = len(recipients)

    unread_count = sum(
        1
        for recipient in recipients
        if not recipient.is_read
    )

    context = {
        "recipients": recipients,
        "total_notifications": total_notifications,
        "unread_count": unread_count,
    }

    return render(
        request,
        "dashboard/notifications.html",
        context
    )

def notification_detail(request, notification_id):
    """
    Display notification details and mark notification as read.
    """

    user_id = request.session.get("yaarlynx_user_id")

    if not user_id:
        return redirect("login")

    user = get_object_or_404(
        UserAccount,
        id=user_id
    )

    recipient = get_object_or_404(
        NotificationRecipient.objects.select_related(
            "notification"
        ),
        notification__notification_id=notification_id,
        user=user,
        notification__is_active=True,
    )

    notification = recipient.notification

    # Check publication status
    if not notification.is_published:

        return redirect("notification_list")

    # Mark as read
    recipient.mark_as_read()

    context = {
        "notification": notification,
        "recipient": recipient,
    }

    return render(
        request,
        "dashboard/notification_detail.html",
        context
    )

def mark_notification_read(request, notification_id):
    """
    Mark a notification as read.
    """

    if request.method != "POST":

        return JsonResponse(
            {
                "success": False,
                "message": "Invalid request method."
            },
            status=400
        )

    user_id = request.session.get(
        "yaarlynx_user_id"
    )

    if not user_id:

        return JsonResponse(
            {
                "success": False,
                "message": "Authentication required."
            },
            status=401
        )

    recipient = get_object_or_404(
        NotificationRecipient,
        notification__notification_id=notification_id,
        user_id=user_id,
    )

    recipient.mark_as_read()

    return JsonResponse(
        {
            "success": True,
            "message": "Notification marked as read."
        }
    )

def mark_all_notifications_read(request):
    """
    Mark all notifications of the logged-in user as read.
    """

    if request.method != "POST":

        return JsonResponse(
            {
                "success": False,
                "message": "Invalid request method."
            },
            status=400
        )

    user_id = request.session.get(
        "yaarlynx_user_id"
    )

    if not user_id:

        return JsonResponse(
            {
                "success": False,
                "message": "Authentication required."
            },
            status=401
        )

    recipients = NotificationRecipient.objects.filter(
        user_id=user_id,
        is_read=False,
        notification__is_active=True,
    )

    updated_count = recipients.update(
        is_read=True,
        read_at=timezone.now(),
    )

    return JsonResponse(
        {
            "success": True,
            "message": "All notifications marked as read.",
            "updated_count": updated_count,
        }
    )

def mark_notification_unread(request, notification_id):
    """
    Mark a notification as unread.
    """

    if request.method != "POST":

        return JsonResponse(
            {
                "success": False,
                "message": "Invalid request method."
            },
            status=400
        )

    user_id = request.session.get(
        "yaarlynx_user_id"
    )

    if not user_id:

        return JsonResponse(
            {
                "success": False,
                "message": "Authentication required."
            },
            status=401
        )

    recipient = get_object_or_404(
        NotificationRecipient,
        notification__notification_id=notification_id,
        user_id=user_id,
    )

    recipient.mark_as_unread()

    return JsonResponse(
        {
            "success": True,
            "message": "Notification marked as unread."
        }
    )




def pdf_overlay_view(request):

    # ========================================================
    # GET ALL OVERLAY IMAGES
    # ========================================================

    overlay_images = OverlayImage.objects.all()


    # ========================================================
    # GET DEFAULT OVERLAY
    # ========================================================

    default_overlay = OverlayImage.objects.filter(
        is_default=True
    ).first()


    # ========================================================
    # POST
    # ========================================================

    if request.method == "POST":

        # ----------------------------------------------------
        # GET UPLOADED PDF
        # ----------------------------------------------------

        pdf_file = request.FILES.get(
            "pdf_file"
        )


        # ----------------------------------------------------
        # Validate PDF
        # ----------------------------------------------------

        if not pdf_file:

            return render(
                request,
                "dashboard/overlay.html",
                {
                    "overlay_images": overlay_images,
                    "default_overlay": default_overlay,
                    "error": "Please upload a PDF file.",
                }
            )


        # ----------------------------------------------------
        # Check extension
        # ----------------------------------------------------

        if not pdf_file.name.lower().endswith(
            ".pdf"
        ):

            return render(
                request,
                "dashboard/overlay.html",
                {
                    "overlay_images": overlay_images,
                    "default_overlay": default_overlay,
                    "error": "Only PDF files are allowed.",
                }
            )


        # ====================================================
        # GET SELECTED OVERLAY
        # ====================================================

        overlay_id = request.POST.get(
            "overlay_image"
        )


        if not overlay_id:

            return render(
                request,
                "dashboard/overlay.html",
                {
                    "overlay_images": overlay_images,
                    "default_overlay": default_overlay,
                    "error": "Please select an overlay image.",
                }
            )


        # ====================================================
        # GET OVERLAY OBJECT
        # ====================================================

        try:

            overlay = OverlayImage.objects.get(
                pk=overlay_id
            )

        except OverlayImage.DoesNotExist:

            return render(
                request,
                "dashboard/overlay.html",
                {
                    "overlay_images": overlay_images,
                    "default_overlay": default_overlay,
                    "error": "Selected overlay image was not found.",
                }
            )


        # ====================================================
        # GET POSITION
        # ====================================================

        position = request.POST.get(
            "position",
            "bottom_center"
        )


        # ====================================================
        # ALLOWED POSITIONS
        # ====================================================

        allowed_positions = {
            "bottom_left",
            "bottom_center",
            "bottom_right",
            "top_left",
            "top_center",
            "top_right",
            "center",
        }


        if position not in allowed_positions:

            position = "bottom_center"


        # ====================================================
        # PROCESS PDF
        # ====================================================

        try:

            result = apply_overlay(
                pdf_file=pdf_file,
                image_path=overlay.image.path,
                position=position,
            )

        except Exception as error:

            return render(
                request,
                "dashboard/overlay.html",
                {
                    "overlay_images": overlay_images,
                    "default_overlay": default_overlay,
                    "error": (
                        f"Unable to process PDF: {error}"
                    ),
                }
            )


        # ====================================================
        # DOWNLOAD
        # ====================================================

        return FileResponse(
            result,
            as_attachment=True,
            filename="pdf_overlay_result.pdf",
            content_type="application/pdf",
        )


    # ========================================================
    # GET REQUEST
    # ========================================================

    return render(
        request,
    "dashboard/overlay.html",
        {
            "overlay_images": overlay_images,
            "default_overlay": default_overlay,
        }
    )